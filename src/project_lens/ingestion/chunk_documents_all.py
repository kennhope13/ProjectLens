"""Create RAG chunks from common office documents, spreadsheets, and images."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pandas as pd
import pytesseract
from bs4 import BeautifulSoup
from langchain_community.document_loaders import CSVLoader, Docx2txtLoader, PyPDFLoader, TextLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from PIL import Image
from pptx import Presentation

from project_lens.paths import DOCUMENTS_DIR, PROCESSED_DIR, PROJECT_ROOT


PROJECT_DIR = PROJECT_ROOT
OUTPUT_FILE = PROCESSED_DIR / "chunks.jsonl"
CHUNK_SIZE = 1_200
CHUNK_OVERLAP = 200

TEXT_EXTENSIONS = {".txt", ".md", ".py", ".log", ".xml", ".rtf"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
SUPPORTED_EXTENSIONS = (
    {".pdf", ".docx", ".pptx", ".xlsx", ".xls", ".csv", ".html", ".htm", ".json"}
    | TEXT_EXTENSIONS
    | IMAGE_EXTENSIONS
)


def load_pdf(path: Path) -> list[Document]:
    return PyPDFLoader(str(path)).load()


def load_text(path: Path) -> list[Document]:
    return TextLoader(str(path), autodetect_encoding=True).load()


def load_csv(path: Path) -> list[Document]:
    return CSVLoader(str(path), autodetect_encoding=True).load()


def load_docx(path: Path) -> list[Document]:
    return Docx2txtLoader(str(path)).load()


def load_pptx(path: Path) -> list[Document]:
    presentation = Presentation(str(path))
    documents = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        text = "\n".join(
            shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)
        )
        if text.strip():
            documents.append(
                Document(
                    page_content=text,
                    metadata={"source": str(path), "slide": slide_number},
                )
            )
    return documents


def load_excel(path: Path) -> list[Document]:
    sheets = pd.read_excel(path, sheet_name=None, dtype=str).items()
    documents = []
    for sheet_name, dataframe in sheets:
        content = dataframe.fillna("").to_csv(index=False)
        if content.strip():
            documents.append(
                Document(
                    page_content=f"Sheet: {sheet_name}\n{content}",
                    metadata={"source": str(path), "sheet_name": str(sheet_name)},
                )
            )
    return documents


def load_html(path: Path) -> list[Document]:
    html = path.read_text(encoding="utf-8", errors="replace")
    return [Document(page_content=BeautifulSoup(html, "html.parser").get_text("\n"), metadata={"source": str(path)})]


def load_json(path: Path) -> list[Document]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    parsed = json.loads(raw)
    return [Document(page_content=json.dumps(parsed, ensure_ascii=False, indent=2), metadata={"source": str(path)})]


def load_image(path: Path) -> list[Document]:
    text = pytesseract.image_to_string(Image.open(path), lang="eng+vie")
    return [Document(page_content=text, metadata={"source": str(path)})] if text.strip() else []


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf(path)
    if suffix == ".docx":
        return load_docx(path)
    if suffix == ".pptx":
        return load_pptx(path)
    if suffix in {".xlsx", ".xls"}:
        return load_excel(path)
    if suffix == ".csv":
        return load_csv(path)
    if suffix in TEXT_EXTENSIONS:
        return load_text(path)
    if suffix in {".html", ".htm"}:
        return load_html(path)
    if suffix == ".json":
        return load_json(path)
    if suffix in IMAGE_EXTENSIONS:
        return load_image(path)
    return []


def create_chunk_id(chunk: Document) -> str:
    payload = json.dumps(
        {"content": chunk.page_content, "metadata": chunk.metadata},
        ensure_ascii=False,
        sort_keys=True,
        default=str,
    )
    return "chunk-" + hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]


def main() -> None:
    if not DOCUMENTS_DIR.exists():
        raise FileNotFoundError(f"Missing documents directory: {DOCUMENTS_DIR}")

    documents: list[Document] = []
    skipped: list[Path] = []
    failed: list[tuple[Path, str]] = []
    for path in DOCUMENTS_DIR.rglob("*"):
        if not path.is_file():
            continue
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            skipped.append(path)
            continue
        try:
            documents.extend(load_file(path))
        except Exception as error:
            failed.append((path, str(error)))

    if not documents:
        raise ValueError("No readable documents found.")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        add_start_index=True,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(documents)
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT_FILE.open("w", encoding="utf-8") as file:
        for chunk in chunks:
            file.write(
                json.dumps(
                    {
                        "id": create_chunk_id(chunk),
                        "page_content": chunk.page_content,
                        "metadata": chunk.metadata,
                    },
                    ensure_ascii=False,
                    default=str,
                )
                + "\n"
            )

    print(f"Loaded {len(documents)} document sections.")
    print(f"Created {len(chunks)} chunks: {OUTPUT_FILE}")
    if skipped:
        print(f"Skipped {len(skipped)} unsupported files.")
    if failed:
        print(f"Failed to read {len(failed)} files:")
        for path, message in failed:
            print(f"- {path.name}: {message}")


if __name__ == "__main__":
    main()
